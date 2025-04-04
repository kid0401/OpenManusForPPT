import asyncio
import os
import shutil
import threading
import time
import tomllib
import uuid
import webbrowser
from datetime import datetime
from functools import partial
from json import dumps
from pathlib import Path

from fastapi import Body, FastAPI, File, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import (
    FileResponse,
    HTMLResponse,
    JSONResponse,
    StreamingResponse,
)
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from pydantic import BaseModel
from pptx import Presentation

app = FastAPI()
app.mount("/static", StaticFiles(directory="static"), name="static")
templates = Jinja2Templates(directory="templates")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# ========== 数据模型 ==========
class Task(BaseModel):
    id: str
    prompt: str
    created_at: datetime
    status: str
    steps: list = []

    def model_dump(self, *args, **kwargs):
        data = super().model_dump(*args, **kwargs)
        data["created_at"] = self.created_at.isoformat()
        return data


# ========== 任务管理器 ==========
class TaskManager:
    def __init__(self):
        self.tasks = {}
        self.queues = {}

    def create_task(self, prompt: str) -> Task:
        task_id = str(uuid.uuid4())
        task = Task(
            id=task_id, prompt=prompt, created_at=datetime.now(), status="pending"
        )
        self.tasks[task_id] = task
        self.queues[task_id] = asyncio.Queue()
        return task

    async def update_task_step(
            self, task_id: str, step: int, result: str, step_type: str = "step"
    ):
        result = str(result) if result is not None else ""
        if task_id in self.tasks:
            task = self.tasks[task_id]
            task.steps.append({"step": step, "result": result, "type": step_type})
            await self.queues[task_id].put(
                {"type": step_type, "step": step, "result": result}
            )
            await self.queues[task_id].put(
                {"type": "status", "status": task.status, "steps": task.steps}
            )

    async def complete_task(self, task_id: str):
        if task_id in self.tasks:
            task = self.tasks[task_id]
            task.status = "completed"
            await self.queues[task_id].put(
                {"type": "status", "status": task.status, "steps": task.steps}
            )
            await self.queues[task_id].put({"type": "complete"})

    async def fail_task(self, task_id: str, error: str):
        if task_id in self.tasks:
            self.tasks[task_id].status = f"failed: {error}"
            await self.queues[task_id].put({"type": "error", "message": error})


# ========== 文件管理器 ==========
class FileManager:
    def __init__(self):
        self.upload_dir = Path("uploads")
        self.upload_dir.mkdir(exist_ok=True)

    async def save_upload_file(self, file: UploadFile) -> Path:
        file_path = self.upload_dir / f"{uuid.uuid4()}-{file.filename}"
        with file_path.open("wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
        return file_path


# ========== 全局实例 ==========
task_manager = TaskManager()
file_manager = FileManager()


@app.post("/api/upload")
async def upload_file(file: UploadFile = File(...)):
    try:
        # 扩展允许的文件类型
        allowed_extensions = {".ppt", ".pptx"}
        allowed_mime_types = {
            "application/vnd.ms-powerpoint",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "application/mspowerpoint",
            "application/octet-stream"
        }

        # 获取文件信息
        file_extension = Path(file.filename).suffix.lower()
        content_type = file.content_type.split(";")[0].strip()

        # 双重验证
        if (file_extension not in allowed_extensions or
                content_type not in allowed_mime_types):
            raise HTTPException(400, "仅支持PPT/PPTX文件")

        if file.size > 5 * 1024 * 1024:
            raise HTTPException(400, "文件大小超过5MB限制")

        # 保存文件
        saved_path = await file_manager.save_upload_file(file)

        # 创建处理任务
        task = task_manager.create_task(f"PPT处理: {file.filename}")
        asyncio.create_task(process_ppt_task(task.id, saved_path))

        return {"task_id": task.id, "filename": file.filename}

    except HTTPException as he:
        raise he
    except Exception as e:
        raise HTTPException(500, f"文件处理失败: {str(e)}")
    finally:
        await file.close()


async def process_ppt_task(task_id: str, file_path: Path):
    try:
        slide_contents = []
        await task_manager.update_task_step(task_id, 1, "正在打开PPT文件...")

        # 读取PPT文件
        prs = Presentation(file_path)

        await task_manager.update_task_step(
            task_id,
            2,
            f"检测到{len(prs.slides)}张幻灯片",
            "analysis"
        )

        # 解析每页幻灯片
        for i, slide in enumerate(prs.slides, 1):
            slide_text = []

            # 提取形状中的文本
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    slide_text.append(shape.text.strip())

            # 过滤空文本并合并
            clean_text = [t for t in slide_text if t]
            content = "\n".join(clean_text) if clean_text else "（空白页）"  # 添加默认值
            slide_contents.append(content)
            await task_manager.update_task_step(
                task_id,
                3,
                f"幻灯片 {i} 内容：\n{content}",
                "slide_content"
            )

            await asyncio.sleep(0.1)  # 避免处理过快

        # 生成总结报告
        slide_contents_str = "\n".join(slide_contents)
        summary = f"PPT解析完成\n总页数：{len(prs.slides)}\n主要包含以下内容：{slide_contents_str}\n"
        print(summary)
        await task_manager.update_task_step(task_id, 4, summary, "summary")


        # 存储分析结果
        result = PPTAnalysisResult(
            task_id=task_id,
            total_slides=len(prs.slides),
            slide_contents=slide_contents,
            summary=summary
        )
        analysis_storage.store_result(task_id, result)

        # 自动触发演讲稿生成
        auto_prompt = (
            "根据PPT内容自动生成演讲稿，要求：\n"
            "1. 包含至少5个具体案例\n"
            "2. 使用口语化表达\n"
            "3. 结构包含引言、主体、结论\n"
            "4. 字数不少于500字\n"
            "5. 语气亲切自然"
        )

        try:
            print("开始自动生成：")
            # 获取PPT分析结果
            ppt_result = analysis_storage.get_result(task_id)
            if not ppt_result:
                raise ValueError("PPT分析结果不存在")

            # 初始化任务状态
            await task_manager.update_task_step(task_id, 1, "开始自动生成演讲稿")

            # 构建生成参数
            generate_params = {
                "task_id": task_id,
                "prompt": auto_prompt
            }

            # 调用生成逻辑
            agent = Manus(
                name="AutoSpeechWriter",
                description="Automatic speech generation based on PPT content"
            )

            # 分阶段生成
            await task_manager.update_task_step(task_id, 2, "分析内容结构")
            outline = await agent.run(
                f"生成演讲稿大纲，要求：{auto_prompt}\nPPT内容：{ppt_result.summary}"
            )

            await task_manager.update_task_step(task_id, 3, "撰写正文内容")
            draft = await agent.run(
                f"根据大纲扩展内容：\n大纲：{outline}\n详细内容：{ppt_result.slide_contents}"
            )

            await task_manager.update_task_step(task_id, 4, "优化表达语气")
            final = await agent.run(
                f"优化以下文本语气：{draft}\n优化要求：使用更口语化的表达，增加过渡语句"
            )

            # 保存结果
            output_dir = Path("output/auto_speeches")
            output_dir.mkdir(parents=True, exist_ok=True)
            output_path = output_dir / f"auto-speech-{task_id}.md"

            with open(output_path, "w", encoding="utf-8") as f:
                f.write(f"# 自动生成演讲稿\n\n{final}")

            # 更新任务状态
            await task_manager.update_task_step(
                task_id,
                5,
                f"自动生成完成：{output_path.name}",
                "auto_generation_done"
            )
            await task_manager.complete_task(task_id)

        except Exception as e:
            await task_manager.fail_task(task_id, f"自动生成失败: {str(e)}")
            raise

        # 保持任务处于活动状态以等待交互
        task_manager.tasks[task_id].status = "awaiting_action"
    except Exception as e:
        await task_manager.fail_task(task_id, f"PPT解析失败: {str(e)}")
    finally:
        file_path.unlink(missing_ok=True)


# ========== 核心路由 ==========
@app.get("/", response_class=HTMLResponse)
async def index(request: Request):
    return templates.TemplateResponse("index.html", {"request": request})

@app.get("/download")
async def download_file(file_path: str):
    if not os.path.exists(file_path):
        raise HTTPException(status_code=404, detail="File not found")

    return FileResponse(file_path, filename=os.path.basename(file_path))


@app.post("/tasks")
async def create_task(prompt: str = Body(..., embed=True)):
    task = task_manager.create_task(prompt)
    asyncio.create_task(run_task(task.id, prompt))
    return {"task_id": task.id}


from app.agent.manus import Manus

async def run_task(task_id: str, prompt: str):
    try:
        task_manager.tasks[task_id].status = "running"

        agent = Manus(
            name="Manus",
            description="A versatile agent that can solve various tasks using multiple tools",
        )

        async def on_think(thought):
            await task_manager.update_task_step(task_id, 0, thought, "think")

        async def on_tool_execute(tool, input):
            await task_manager.update_task_step(
                task_id, 0, f"Executing tool: {tool}\nInput: {input}", "tool"
            )

        async def on_action(action):
            await task_manager.update_task_step(
                task_id, 0, f"Executing action: {action}", "act"
            )

        async def on_run(step, result):
            await task_manager.update_task_step(task_id, step, result, "run")

        from app.logger import logger

        class SSELogHandler:
            def __init__(self, task_id):
                self.task_id = task_id

            async def __call__(self, message):
                import re

                # Extract - Subsequent Content
                cleaned_message = re.sub(r"^.*? - ", "", message)

                event_type = "log"
                if "✨ Manus's thoughts:" in cleaned_message:
                    event_type = "think"
                elif "🛠️ Manus selected" in cleaned_message:
                    event_type = "tool"
                elif "🎯 Tool" in cleaned_message:
                    event_type = "act"
                elif "📝 Oops!" in cleaned_message:
                    event_type = "error"
                elif "🏁 Special tool" in cleaned_message:
                    event_type = "complete"

                await task_manager.update_task_step(
                    self.task_id, 0, cleaned_message, event_type
                )

        sse_handler = SSELogHandler(task_id)
        logger.add(sse_handler)

        result = await agent.run(prompt)
        await task_manager.update_task_step(task_id, 1, result, "result")
        await task_manager.complete_task(task_id)
    except Exception as e:
        await task_manager.fail_task(task_id, str(e))


@app.get("/tasks/{task_id}/events")
async def task_events(task_id: str):
    async def event_generator():
        if task_id not in task_manager.queues:
            yield f"event: error\ndata: {dumps({'message': 'Task not found'})}\n\n"
            return

        queue = task_manager.queues[task_id]

        task = task_manager.tasks.get(task_id)
        if task:
            yield f"event: status\ndata: {dumps({'type': 'status', 'status': task.status, 'steps': task.steps})}\n\n"

        while True:
            try:
                event = await queue.get()
                formatted_event = dumps(event)

                yield ": heartbeat\n\n"

                if event["type"] == "complete":
                    yield f"event: complete\ndata: {formatted_event}\n\n"
                    break
                elif event["type"] == "error":
                    yield f"event: error\ndata: {formatted_event}\n\n"
                    break
                elif event["type"] == "step":
                    task = task_manager.tasks.get(task_id)
                    if task:
                        yield f"event: status\ndata: {dumps({'type': 'status', 'status': task.status, 'steps': task.steps})}\n\n"
                    yield f"event: {event['type']}\ndata: {formatted_event}\n\n"
                elif event["type"] in ["think", "tool", "act", "run"]:
                    yield f"event: {event['type']}\ndata: {formatted_event}\n\n"
                else:
                    yield f"event: {event['type']}\ndata: {formatted_event}\n\n"

            except asyncio.CancelledError:
                print(f"Client disconnected for task {task_id}")
                break
            except Exception as e:
                print(f"Error in event stream: {str(e)}")
                yield f"event: error\ndata: {dumps({'message': str(e)})}\n\n"
                break

    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        },
    )


@app.get("/tasks")
async def get_tasks():
    sorted_tasks = sorted(
        task_manager.tasks.values(), key=lambda task: task.created_at, reverse=True
    )
    return JSONResponse(
        content=[task.model_dump() for task in sorted_tasks],
        headers={"Content-Type": "application/json"},
    )


@app.get("/tasks/{task_id}")
async def get_task(task_id: str):
    if task_id not in task_manager.tasks:
        raise HTTPException(status_code=404, detail="Task not found")
    return task_manager.tasks[task_id]


@app.get("/config/status")
async def check_config_status():
    config_path = Path(__file__).parent / "config" / "config.toml"
    example_config_path = Path(__file__).parent / "config" / "config.example.toml"

    if config_path.exists():
        try:
            with open(config_path, "rb") as f:
                current_config = tomllib.load(f)
            return {"status": "exists", "config": current_config}
        except Exception as e:
            return {"status": "error", "message": str(e)}
    elif example_config_path.exists():
        try:
            with open(example_config_path, "rb") as f:
                example_config = tomllib.load(f)
            return {"status": "missing", "example_config": example_config}
        except Exception as e:
            return {"status": "error", "message": str(e)}
    else:
        return {"status": "no_example"}


@app.post("/config/save")
async def save_config(config_data: dict = Body(...)):
    try:
        config_dir = Path(__file__).parent / "config"
        config_dir.mkdir(exist_ok=True)

        config_path = config_dir / "config.toml"

        toml_content = ""

        if "llm" in config_data:
            toml_content += "# Global LLM configuration\n[llm]\n"
            llm_config = config_data["llm"]
            for key, value in llm_config.items():
                if key != "vision":
                    if isinstance(value, str):
                        toml_content += f'{key} = "{value}"\n'
                    else:
                        toml_content += f"{key} = {value}\n"

        if "server" in config_data:
            toml_content += "\n# Server configuration\n[server]\n"
            server_config = config_data["server"]
            for key, value in server_config.items():
                if isinstance(value, str):
                    toml_content += f'{key} = "{value}"\n'
                else:
                    toml_content += f"{key} = {value}\n"

        with open(config_path, "w", encoding="utf-8") as f:
            f.write(toml_content)

        return {"status": "success"}
    except Exception as e:
        return {"status": "error", "message": str(e)}


@app.exception_handler(Exception)
async def generic_exception_handler(request: Request, exc: Exception):
    return JSONResponse(
        status_code=500, content={"message": f"Server error: {str(exc)}"}
    )


def open_local_browser(config):
    webbrowser.open_new_tab(f"http://{config['host']}:{config['port']}")


def load_config():
    try:
        config_path = Path(__file__).parent / "config" / "config.toml"

        if not config_path.exists():
            return {"host": "localhost", "port": 5172}

        with open(config_path, "rb") as f:
            config = tomllib.load(f)

        return {"host": config["server"]["host"], "port": config["server"]["port"]}
    except FileNotFoundError:
        return {"host": "localhost", "port": 5172}
    except KeyError as e:
        print(
            f"The configuration file is missing necessary fields: {str(e)}, use default configuration"
        )
        return {"host": "localhost", "port": 5172}


# ========== 新增数据模型 ==========
class PPTAnalysisResult(BaseModel):
    task_id: str
    total_slides: int
    slide_contents: list[str]
    summary: str


# ========== 新增存储服务 ==========
class AnalysisStorage:
    def __init__(self):
        self.results = {}

    def store_result(self, task_id: str, result: PPTAnalysisResult):
        self.results[task_id] = result

    def get_result(self, task_id: str) -> PPTAnalysisResult:
        return self.results.get(task_id)

analysis_storage = AnalysisStorage()

@app.post("/api/generate-speech")
async def generate_speech(task_id: str = Body(...), prompt: str = Body(...)):
    try:
        # 获取PPT分析结果
        ppt_result = analysis_storage.get_result(task_id)
        if not ppt_result:
            raise HTTPException(404, "PPT analysis result not found")

        # 创建新任务
        new_task = task_manager.create_task(f"生成演讲稿 - {prompt[:20]}...")

        async def speech_generation_process():
            try:
                # 获取Manus agent实例
                agent = Manus(
                    name="SpeechWriter",
                    description="Specialized in generating speeches based on PPT content"
                )

                # 构建详细提示
                full_prompt = f"根据以下PPT内容：\n{ppt_result.summary}\n\n幻灯片详细内容：\n" + \
                              "\n\n".join(
                                  [f"第{i + 1}页：{content}" for i, content in enumerate(ppt_result.slide_contents)]) + \
                              f"\n\n用户要求：{prompt}"

                # 分步骤处理
                await task_manager.update_task_step(new_task.id, 1, "正在分析PPT内容...")
                await asyncio.sleep(0.5)

                await task_manager.update_task_step(new_task.id, 2, "构建演讲框架...")
                await asyncio.sleep(0.5)

                await task_manager.update_task_step(new_task.id, 3, "生成演讲内容...")

                # 调用agent生成内容
                result = await agent.run(full_prompt)

                # 保存结果
                output_dir = Path("output/speeches")
                output_dir.mkdir(parents=True, exist_ok=True)
                output_path = output_dir / f"speech-{task_id}.md"

                with open(output_path, "w", encoding="utf-8") as f:
                    f.write(result)

                await task_manager.update_task_step(
                    new_task.id,
                    4,
                    f"演讲稿生成完成：{output_path.name}",
                    "file_generated"
                )

                await task_manager.complete_task(new_task.id)

                return {"path": str(output_path)}

            except Exception as e:
                await task_manager.fail_task(new_task.id, str(e))
                raise

        asyncio.create_task(speech_generation_process())
        return {"task_id": new_task.id}

    except Exception as e:
        raise HTTPException(500, f"生成失败: {str(e)}")


@app.get("/api/download-speech/{task_id}")
async def download_speech(task_id: str):
    file_path = Path(f"output/speeches/speech-{task_id}.md")
    if not file_path.exists():
        raise HTTPException(404, "File not found")

    return FileResponse(file_path, filename=file_path.name)

if __name__ == "__main__":
    import uvicorn

    config = load_config()
    open_with_config = partial(open_local_browser, config)
    threading.Timer(3, open_with_config).start()
    uvicorn.run(app, host=config["host"], port=config["port"])